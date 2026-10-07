import io
import os
import datetime
from flask import abort, current_app, render_template, redirect, send_file, url_for, flash, request, jsonify, send_from_directory
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from sqlalchemy import or_

from app import db
from app.main import bp
from app.models import ChatMessage, Document, DocumentApproval, Evaluation, Lecturer, Student, Seminar, Supervision, Bimbingan, SupervisorApplication, User, Honorarium

# Impor TAHAPAN_TUGAS dari student.routes untuk konsistensi
try:
    from app.student.routes import TAHAPAN_TUGAS
except ImportError:
    TAHAPAN_TUGAS = {}

ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}

CRITERIA_DEFINITIONS = {
    'Seminar Proposal': {
        "BAB I Pendahuluan": [
            {'id': 'bab1_a', 'label': 'A. Latar Belakang'},
            {'id': 'bab1_b', 'label': 'B. Rumusan Masalah'},
            {'id': 'bab1_c', 'label': 'C. Maksud dan Tujuan Penelitian'},
            {'id': 'bab1_d', 'label': 'D. Kegunaan Penelitian'}
        ],
        "BAB II Tinjauan Pustaka": [
            {'id': 'bab2_a', 'label': 'A. Tinjauan Penelitian Terkait'},
            {'id': 'bab2_b', 'label': 'B. Konsep dan Teori yang Digunakan'},
            {'id': 'bab2_c', 'label': 'C. Kerangka Pemikiran'},
            {'id': 'bab2_d', 'label': 'D. Hipotesis (Jika Ada)'}
        ],
        "BAB III Metode Penelitian": [
            {'id': 'bab3_a', 'label': 'A. Pendekatan Yang Digunakan'},
            {'id': 'bab3_b', 'label': 'B. Populasi dan Sampel'},
            {'id': 'bab3_c', 'label': 'C. Unit Analisis'},
            {'id': 'bab3_d', 'label': 'D. Analisis Instrumen'},
            {'id': 'bab3_e', 'label': 'E. Definisi Operasional'},
            {'id': 'bab3_f', 'label': 'F. Instrumen/Pedoman Wawancara'}
        ],
        "Daftar Pustaka": [
            {'id': 'pustaka_sumber', 'label': 'Sumber'},
            {'id': 'pustaka_format', 'label': 'Format Penulisan'},
            {'id': 'pustaka_kutipan', 'label': 'Pengutipan Sumber'},
            {'id': 'pustaka_plagiarisme', 'label': 'Aspek Plagiarisme'}
        ],
        "Performance": [
            {'id': 'perf_gagasan', 'label': 'Kemampuan Mengemukakan Gagasan'},
            {'id': 'perf_argumen', 'label': 'Kemampuan Berargumentasi'},
            {'id': 'perf_ilmiah', 'label': 'Sikap Ilmiah'},
            {'id': 'perf_personal', 'label': 'Sikap Personal dan Sosial'}
        ]
    },
    'Seminar Hasil': {
        "BAB IV Hasil dan Pembahasan": [
            {'id': 'bab4_deskripsi', 'label': 'A. Deskripsi Data dan Hasil'},
            {'id': 'bab4_temuan', 'label': 'B. Temuan Hasil Penelitian'},
            {'id': 'bab4_pembahasan', 'label': 'C. Pembahasan Hasil Penelitian'}
        ],
        "BAB V Kesimpulan dan Saran": [
            {'id': 'bab5_kesimpulan', 'label': 'A. Kesimpulan'},
            {'id': 'bab5_saran', 'label': 'B. Keterbatasan dan Saran'}
        ],
        "Performance": [
            {'id': 'perf_gagasan', 'label': 'Kemampuan Mengemukakan Gagasan'},
            {'id': 'perf_argumen', 'label': 'Kemampuan Berargumentasi'},
            {'id': 'perf_ilmiah', 'label': 'Sikap Ilmiah'},
            {'id': 'perf_personal', 'label': 'Sikap Personal dan Sosial'}
        ]
    },
    'Sidang Tertutup': { 
        "BAB I Pendahuluan": [{'id': 'sidang_bab1', 'label': 'Kualitas Latar Belakang, Rumusan Masalah, dan Tujuan'}],
        "BAB II Tinjauan Pustaka": [{'id': 'sidang_bab2', 'label': 'Relevansi, Kedalaman, dan Kemutakhiran Teori'}],
        "BAB III Metode Penelitian": [{'id': 'sidang_bab3', 'label': 'Ketepatan Metodologi, Desain, dan Teknik Analisis'}],
        "BAB IV Hasil dan Pembahasan": [{'id': 'sidang_bab4', 'label': 'Penyajian, Analisis, dan Pembahasan Hasil'}],
        "BAB V Kesimpulan dan Saran": [{'id': 'sidang_bab5', 'label': 'Konsistensi Kesimpulan, Implikasi, dan Saran'}],
        "Performance Sidang": [
            {'id': 'sidang_perf_presentasi', 'label': 'Kualitas Presentasi (Kejelasan, Sistematika, Waktu)'},
            {'id': 'sidang_perf_penguasaan', 'label': 'Penguasaan Materi Keseluruhan'},
            {'id': 'sidang_perf_argumentasi', 'label': 'Kemampuan Berargumentasi dan Menjawab'}
        ]
    },
    'Sidang Terbuka': {
        "BAB I Pendahuluan": [{'id': 'sidang_bab1', 'label': 'Kualitas Latar Belakang, Rumusan Masalah, dan Tujuan'}],
        "BAB II Tinjauan Pustaka": [{'id': 'sidang_bab2', 'label': 'Relevansi, Kedalaman, dan Kemutakhiran Teori'}],
        "BAB III Metode Penelitian": [{'id': 'sidang_bab3', 'label': 'Ketepatan Metodologi, Desain, dan Teknik Analisis'}],
        "BAB IV Hasil dan Pembahasan": [{'id': 'sidang_bab4', 'label': 'Penyajian, Analisis, dan Pembahasan Hasil'}],
        "BAB V Kesimpulan dan Saran": [{'id': 'sidang_bab5', 'label': 'Konsistensi Kesimpulan, Implikasi, dan Saran'}],
        "Performance Sidang": [
            {'id': 'sidang_perf_presentasi', 'label': 'Kualitas Presentasi (Kejelasan, Sistematika, Waktu)'},
            {'id': 'sidang_perf_penguasaan', 'label': 'Penguasaan Materi Keseluruhan'},
            {'id': 'sidang_perf_argumentasi', 'label': 'Kemampuan Berargumentasi dan Menjawab'}
        ]
    },
    'Sidang Komisi Pra-Seminar Proposal': {
        "On Desk": [
            {'id': 'pp_desk_1', 'label': '1. Pendahuluan'},
            {'id': 'pp_desk_2', 'label': '2. Tinjauan Pustaka'},
            {'id': 'pp_desk_3', 'label': '3. Metode'},
            {'id': 'pp_desk_4', 'label': '4. Instrumen Penelitian'},
            {'id': 'pp_desk_5', 'label': '5. Daftar Pustaka'}
        ],
        "Performance": [
            {'id': 'pp_perf_1', 'label': '1. Kemampuan Mengemukakan Gagasan Inti'},
            {'id': 'pp_perf_2', 'label': '2. Kemampuan Berargumen'},
            {'id': 'pp_perf_3', 'label': '3. Sikap Ilmiah'},
            {'id': 'pp_perf_4', 'label': '4. Sikap Personal dan Sosial'}
        ]
    },
    'Sidang Komisi Pra-Seminar Hasil': {
         "On Desk": [
            {'id': 'psh_desk_1', 'label': '1. Pendahuluan'},
            {'id': 'psh_desk_2', 'label': '2. Tinjauan Pustaka'},
            {'id': 'psh_desk_3', 'label': '3. Validasi Instrumen Penelitian'},
            {'id': 'psh_desk_4', 'label': '4. Temuan hasil penelitian'},
            {'id': 'psh_desk_5', 'label': '5. Daftar Pustaka'}
        ],
        "Performance": [
            {'id': 'psh_perf_1', 'label': '1. Kemampuan Mengemukakan Gagasan Inti'},
            {'id': 'psh_perf_2', 'label': '2. Kemampuan Berargumen'},
            {'id': 'psh_perf_3', 'label': '3. Sikap Ilmiah'},
            {'id': 'psh_perf_4', 'label': '4. Sikap Personal dan Sosial'}
        ]
    },
    'Sidang Komisi Pra-Ujian Tertutup': {
        "On Desk": [
            {'id': 'put_desk_1', 'label': '1. Pendahuluan'},
            {'id': 'put_desk_2', 'label': '2. Tinjauan Pustaka'},
            {'id': 'put_desk_3', 'label': '3. Metode'},
            {'id': 'put_desk_4', 'label': '4. Instrumen Penelitian'},
            {'id': 'put_desk_5', 'label': '5. Kesimpulan dan saran'},
            {'id': 'put_desk_6', 'label': '6. Daftar Pustaka'}
        ],
        "Performance": [
            {'id': 'put_perf_1', 'label': '1. Kemampuan Mengemukakan Gagasan Inti'},
            {'id': 'put_perf_2', 'label': '2. Kemampuan Berargumen'},
            {'id': 'put_perf_3', 'label': '3. Sikap Ilmiah'},
            {'id': 'put_perf_4', 'label': '4. Sikap Personal dan Sosial'}
        ]
    }
}

def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@bp.route('/uploads/<path:filename>')
@login_required
def serve_upload(filename):
    return send_from_directory(current_app.config['UPLOAD_FOLDER'], filename)

@bp.route('/')
@login_required
def index():
    if current_user.role == 'dosen':
        return redirect(url_for('main.dashboard'))
    elif current_user.role == 'admin':
        return redirect(url_for('main.kelola_panduan')) 
    elif current_user.role == 'mahasiswa':
        return redirect(url_for('student.dashboard'))
    else:
        abort(403)
        
@bp.route('/dashboard')
@login_required
def dashboard():
    if current_user.role != 'dosen':
        abort(403)
    
    lecturer = current_user.lecturer
    my_supervisions = lecturer.supervisions 
    
    stats = {
        'mahasiswa_dibimbing': len(my_supervisions),
        'seminar_proposal': 0,
        'seminar_hasil': 0,
        'sidang_tertutup': 0,
        'sidang_terbuka': 0,
        'yudisium': 0
    }

    for sup in my_supervisions:
        if not sup.student: continue
        step = sup.student.current_progress_step
        if step == 3: stats['seminar_proposal'] += 1
        elif step == 6: stats['seminar_hasil'] += 1
        elif step == 9: stats['sidang_tertutup'] += 1
        elif step == 11: stats['sidang_terbuka'] += 1
        elif step >= 13: stats['yudisium'] += 1

    student_data = []
    
    for supervision in my_supervisions:
        student = supervision.student
        if not student: continue

        def get_name_by_role(target_student, role_name):
            for s in target_student.supervisions:
                if s.role == role_name:
                    return s.lecturer.name if s.lecturer else '-'
            return '-'

        student_data.append({
            'student': student,
            'my_role': supervision.role,
            'promotor_name': get_name_by_role(student, 'Promotor'),
            'copro1_name': get_name_by_role(student, 'Ko-Promotor 1'),
            'copro2_name': get_name_by_role(student, 'Ko-Promotor 2')
        })

    upcoming_seminars = Seminar.query.join(Student).join(Supervision).filter(
        Supervision.lecturer_id == lecturer.id,
        Seminar.date >= datetime.date.today()
    ).order_by(Seminar.date.asc()).all()

    return render_template('main/dashboard.html', 
                           title='Dashboard Dosen', 
                           stats=stats, 
                           student_data=student_data,
                           upcoming_seminars=upcoming_seminars)

@bp.route('/pembimbingan/detail/<int:student_id>')
@login_required
def pembimbingan_detail(student_id):
    if current_user.role != 'dosen': abort(403)
    
    student = Student.query.get_or_404(student_id)
    is_supervisor = any(s.lecturer_id == current_user.lecturer.id for s in student.supervisions)
    if not is_supervisor: abort(403)

    required_tasks_info = TAHAPAN_TUGAS.get(student.current_progress_step, [])
    docs_this_step = student.documents.filter_by(progress_step=student.current_progress_step).all()
    uploaded_docs = {doc.document_type: doc for doc in docs_this_step}

    tasks_with_status = []
    for task_info in required_tasks_info:
        doc_type = task_info['type']
        uploaded_document = uploaded_docs.get(doc_type)
        task = {
            'type': doc_type, 'icon': task_info['icon'], 'uploaded_doc': uploaded_document
        }
        tasks_with_status.append(task)
    
    approved_doc_ids = []
    if current_user.lecturer:
        approvals = DocumentApproval.query.filter_by(lecturer_id=current_user.lecturer.id, is_approved=True).all()
        approved_doc_ids = [a.document_id for a in approvals]

    return render_template('main/pembimbingan_detail.html',
                           title=f'Bimbingan: {student.name}',
                           student=student,
                           tasks_with_status=tasks_with_status,
                           approved_doc_ids=approved_doc_ids)

@bp.route('/daftar-ajuan')
@login_required
def daftar_ajuan():
    if current_user.role != 'dosen': abort(403)
    ajuan_list = Bimbingan.query.filter_by(lecturer_id=current_user.lecturer.id).order_by(Bimbingan.created_at.desc()).all()
    return render_template('main/daftar_ajuan_mahasiswa.html', ajuan_list=ajuan_list)

@bp.route('/update_ajuan_status', methods=['POST'])
@login_required
def update_ajuan_status():
    if current_user.role != 'dosen': abort(403)
    data = request.get_json()
    ajuan = Bimbingan.query.get_or_404(data['ajuan_id'])
    
    if ajuan.lecturer_id != current_user.lecturer.id:
        return jsonify({'status': 'error', 'message': 'Unauthorized'}), 403
        
    ajuan.status = data['status']
    db.session.commit()
    return jsonify({'status': 'success', 'message': 'Status berhasil diperbarui.', 'new_status': ajuan.status})

@bp.route('/masukan_hasil_seminar')
@login_required
def masukan_hasil_seminar():
    if current_user.role != 'dosen': abort(403)
    student_ids = [s.student_id for s in current_user.lecturer.supervisions]
    seminars = Seminar.query.filter(Seminar.student_id.in_(student_ids)).order_by(Seminar.date.desc()).all()
    evaluated_seminar_ids = []
    evals = Evaluation.query.filter_by(evaluator_id=current_user.lecturer.id).all()
    evaluated_seminar_ids = [e.seminar_id for e in evals]
    return render_template('main/masukan_hasil_seminar.html', seminars=seminars, evaluated_seminar_ids=evaluated_seminar_ids)

@bp.route('/submit_seminar_grade', methods=['POST'])
@login_required
def submit_seminar_grade():
    if current_user.role != 'dosen': return jsonify({'status': 'error', 'message': 'Unauthorized'}), 403
    data = request.get_json()
    seminar_id = data.get('seminar_id')
    final_grade = data.get('final_grade')
    general_notes = data.get('general_notes')
    criteria_scores = data.get('criteria_scores')
    
    seminar = Seminar.query.get_or_404(seminar_id)
    evaluation = Evaluation(
        seminar_id=seminar.id,
        student_id=seminar.student_id,
        evaluator_id=current_user.lecturer.id,
        final_grade=final_grade,
        general_notes=general_notes,
        criteria_scores=criteria_scores
    )
    db.session.add(evaluation)
    seminar.final_grade = final_grade 
    db.session.commit()
    return jsonify({'status': 'success', 'message': 'Nilai berhasil disimpan.'})

@bp.route('/laporan/progres')
@login_required
def laporan_progres():
    if current_user.role != 'dosen': abort(403)
    students = [s.student for s in current_user.lecturer.supervisions]
    return render_template('main/laporan_progres.html', students=students)

@bp.route('/laporan/honorarium')
@login_required
def laporan_honorarium():
    if current_user.role != 'dosen': abort(403)
    honorariums = Honorarium.query.filter_by(lecturer_id=current_user.lecturer.id).all()
    return render_template('main/laporan_honorarium.html', honorariums=honorariums)

@bp.route('/laporan/akumulasi')
@login_required
def akumulasi_penilaian_tapd():
    if current_user.role != 'dosen': abort(403)
    students = [s.student for s in current_user.lecturer.supervisions]
    return render_template('main/akumulasi_penilaian_tapd.html', students=students)

@bp.route('/api/student/<int:student_id>/seminars')
@login_required
def get_student_seminars_api(student_id):
    student = Student.query.get_or_404(student_id)
    seminars = Seminar.query.filter_by(student_id=student.id).all()
    data = [{'id': s.id, 'type': s.seminar_type, 'date': s.date.strftime('%Y-%m-%d')} for s in seminars]
    return jsonify(data)

@bp.route('/api/student/<int:student_id>/zoom-link', methods=['POST'])
@login_required
def update_zoom_link(student_id):
    if current_user.role != 'dosen': return jsonify({'status': 'error', 'message': 'Unauthorized'}), 403
    student = Student.query.get_or_404(student_id)
    is_supervisor = any(s.lecturer_id == current_user.lecturer.id for s in student.supervisions)
    if not is_supervisor: return jsonify({'status': 'error', 'message': 'Anda bukan pembimbing mahasiswa ini'}), 403
        
    data = request.get_json()
    zoom_link = data.get('zoomLink')
    if not zoom_link: return jsonify({'status': 'error', 'message': 'Link Zoom tidak boleh kosong'}), 400
        
    student.zoom_link = zoom_link
    db.session.commit()
    return jsonify({'status': 'success', 'message': 'Link Zoom berhasil disimpan'})

@bp.route('/api/student/<int:student_id>/submit-sidang-grade', methods=['POST'])
@login_required
def submit_sidang_grade(student_id):
    if current_user.role != 'dosen': abort(403)
    data = request.get_json()
    grade = data.get('grade')
    notes = data.get('notes')
    if not grade: return jsonify({'status': 'error', 'message': 'Nilai harus diisi.'}), 400
    student = Student.query.get_or_404(student_id)
    student.sidang_grade = grade
    student.sidang_notes = notes
    db.session.commit()
    return jsonify({'status': 'success', 'message': 'Nilai sidang berhasil disimpan.'})

@bp.route('/export-evaluations', methods=['POST'])
@login_required
def export_evaluations():
    if current_user.role != 'dosen':
        abort(403)

    data = request.get_json()
    seminar_ids = data.get('seminar_ids')
    if not seminar_ids:
        return jsonify({'status': 'error', 'message': 'Tidak ada seminar yang dipilih.'}), 400

    output = io.BytesIO()
    workbook = Workbook()
    workbook.remove(workbook.active) 

    for seminar_id in seminar_ids:
        seminar = Seminar.query.get(seminar_id)
        if not seminar: continue

        student = seminar.student_profile
        evaluations = Evaluation.query.filter_by(seminar_id=seminar.id).all()
        
        sheet_title = seminar.seminar_type.replace('Pra-', 'Pra ').title()
        sheet = workbook.create_sheet(title=sheet_title[:31]) 

        # Styling
        bold_font = Font(bold=True)
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        grey_fill = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")

        # Header Mahasiswa
        sheet.merge_cells('A1:B1'); sheet['A1'] = 'NAMA MAHASISWA'; sheet['A1'].font = bold_font
        sheet.merge_cells('C1:G1'); sheet['C1'] = f': {student.name}'
        sheet.merge_cells('A2:B2'); sheet['A2'] = 'NIM'; sheet['A2'].font = bold_font
        sheet.merge_cells('C2:G2'); sheet['C2'] = f': {student.nim}'
        sheet.merge_cells('A3:B3'); sheet['A3'] = 'JUDUL'; sheet['A3'].font = bold_font
        sheet.merge_cells('C3:G3'); sheet['C3'] = f': {student.research_title}'
        sheet.column_dimensions['A'].width = 5
        sheet.column_dimensions['B'].width = 35
        
        header_row = 5
        sheet.cell(row=header_row, column=1, value="NO").font = bold_font
        sheet.cell(row=header_row, column=2, value="KOMPONEN PENILAIAN").font = bold_font
        
        evaluators = [Lecturer.query.get(e.evaluator_id) for e in evaluations]
        for i, evaluator in enumerate(evaluators, 0):
            col = 3 + i
            sheet.cell(row=header_row, column=col, value=evaluator.name).font = bold_font
            sheet.cell(row=header_row, column=col).alignment = center_align
            sheet.column_dimensions[chr(67 + i)].width = 20 # 67 adalah ASCII untuk 'C'
        
        current_row = header_row + 1
        criteria_structure = CRITERIA_DEFINITIONS.get(seminar.seminar_type, {})
        evaluator_total_scores = {ev.id: 0 for ev in evaluators}
        
        item_number = 1
        for section, criteria_list in criteria_structure.items():
            cell = sheet.cell(row=current_row, column=1, value=section)
            cell.font = bold_font; cell.fill = grey_fill
            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2 + len(evaluators))
            current_row += 1
            
            for criterion in criteria_list:
                sheet.cell(row=current_row, column=1, value=item_number)
                sheet.cell(row=current_row, column=2, value=criterion['label'])
                item_number += 1
                
                for i, evaluator in enumerate(evaluators, 0):
                    col = 3 + i
                    evaluation = next((e for e in evaluations if e.evaluator_id == evaluator.id), None)
                    score = 0
                    if evaluation and evaluation.criteria_scores:
                        score_str = evaluation.criteria_scores.get(f"{criterion['id']}_score", "0")
                        score = int(score_str) if score_str and score_str.isdigit() else 0
                        evaluator_total_scores[evaluator.id] += score
                    sheet.cell(row=current_row, column=col, value=score).alignment = center_align
                current_row += 1

        # Summary Section 
        current_row += 1 
        
        num_components = sum(len(v) for v in criteria_structure.values())
        total_overall_average = 0

        sheet.cell(row=current_row, column=2, value="TOTAL NILAI").font = bold_font
        for i, evaluator in enumerate(evaluators, 0):
            col = 3 + i
            sheet.cell(row=current_row, column=col, value=evaluator_total_scores[evaluator.id]).alignment = center_align
        current_row += 1

        sheet.cell(row=current_row, column=2, value="RATA-RATA NILAI").font = bold_font
        for i, evaluator in enumerate(evaluators, 0):
            col = 3 + i
            avg_score = (evaluator_total_scores[evaluator.id] / num_components) if num_components > 0 else 0
            sheet.cell(row=current_row, column=col, value=round(avg_score, 2)).alignment = center_align
            total_overall_average += avg_score
        current_row += 2 

        final_average = (total_overall_average / len(evaluators)) if evaluators else 0
        grade, gpa_range, status = ('', '', '')
        if final_average >= 85: grade, gpa_range, status = 'A', '3,51 - 4,00', 'Lulus'
        elif final_average >= 80: grade, gpa_range, status = 'A-', '3,25 - 3,50', 'Lulus'
        elif final_average >= 75: grade, gpa_range, status = 'B', '3,00 - 3,24', 'Tidak Lulus'
        else: grade, gpa_range, status = 'C', '< 3,00', 'Tidak Lulus'

        summary_rows = [("NILAI AKHIR", round(final_average, 2)), ("NILAI KONVERSI", gpa_range), ("GRADE", grade)]
        for label, value in summary_rows:
            sheet.cell(row=current_row, column=4, value=label).font = bold_font
            sheet.cell(row=current_row, column=5, value=value)
            current_row += 1

        sheet.cell(row=current_row, column=1, value="KETERANGAN HASIL SIDANG").font = bold_font
        sheet.cell(row=current_row, column=5, value=status).font = bold_font
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        sheet.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=6)

    workbook.save(output)
    output.seek(0)
    
    return send_file(
        output, as_attachment=True,
        download_name=f'Laporan Nilai Lengkap_{student.name.replace(" ", "_")}_{student.nim}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@bp.route('/api/get-proposal-evaluation/<int:seminar_hasil_id>')
@login_required
def get_proposal_evaluation(seminar_hasil_id):
    if current_user.role != 'dosen':
        abort(403)

    seminar_hasil = Seminar.query.get_or_404(seminar_hasil_id)
    student_id = seminar_hasil.student_id

    seminar_proposal = Seminar.query.filter_by(
        student_id=student_id, 
        seminar_type='Seminar Proposal'
    ).first()

    if not seminar_proposal:
        return jsonify({'error': 'Seminar Proposal not found for this student'}), 404

    proposal_evaluation = Evaluation.query.filter_by(seminar_id=seminar_proposal.id).first()

    if proposal_evaluation and proposal_evaluation.criteria_scores:
        return jsonify(proposal_evaluation.criteria_scores)
    
    return jsonify({})

@bp.route('/admin/kelola-panduan', methods=['GET', 'POST'])
@login_required
def kelola_panduan():
    if current_user.role != 'admin':
        abort(403)
        
    panduan_dir = os.path.join(current_app.root_path, 'static', 'panduan')
    os.makedirs(panduan_dir, exist_ok=True)
    
    if request.method == 'POST':
        # 1. Handle Upload Dokumen Panduan Umum (Gabungan 4 File Lama + 6 File Baru + 1 Dosen)
        file_mapping = {
            'tautan_1': 'Cara_Pembayaran_Wisuda_dan_Pembuatan_Poster.pdf',
            'dokumen_yudisium': 'Dokumen_Perlengkapan_Yudisium.pdf', # File lama
            'pakta_integritas': 'Pakta_Integritas_Publish_Paper.pdf', # File lama
            'panduan_sidang': 'Panduan_Pelaksanaan_Sidang.pdf', # File lama
            'panduan_proposal': 'Panduan_Proposal_dan_TAPD.pdf', # File lama
            'tautan_6': 'Panduan_Tugas_Literature_Review.pdf',
            'tautan_7': 'RPS_MK_Penulisan_Proposal.pdf',
            'tautan_8': 'Response_Form.pdf',
            'tautan_9': 'Slide_Materi_MK_Penulisan_Proposal.pdf',
            'tautan_10': 'Surat_Keterangan_Index_Publikasi.pdf',
            'panduan_dosen': 'Panduan_Dosen_Pembimbing.pdf' # File dosen lama
        }
        
        for input_name, filename in file_mapping.items():
            f = request.files.get(input_name)
            if f and f.filename != '':
                if f.filename.endswith('.pdf'):
                    f.save(os.path.join(panduan_dir, filename))
                else:
                    flash(f'File untuk {input_name} harus berformat PDF.', 'danger')

        # 2. Handle Upload Template Persetujuan Dinamis per Tahap
        tahap = request.form.get('template_tahap')
        file_tpl = request.files.get('file_template')
        
        if tahap and file_tpl and file_tpl.filename != '':
            ext = file_tpl.filename.rsplit('.', 1)[1].lower()
            if ext in ['doc', 'docx']:
                filename = f'Template_Persetujuan_Tahap_{tahap}.{ext}'
                
                old_doc = os.path.join(panduan_dir, f'Template_Persetujuan_Tahap_{tahap}.doc')
                old_docx = os.path.join(panduan_dir, f'Template_Persetujuan_Tahap_{tahap}.docx')
                if os.path.exists(old_doc): os.remove(old_doc)
                if os.path.exists(old_docx): os.remove(old_docx)
                
                file_tpl.save(os.path.join(panduan_dir, filename))
                flash(f'Template Persetujuan untuk Tahap {tahap} berhasil diperbarui.', 'success')
            else:
                flash('Template Persetujuan harus berformat Word (.doc/.docx).', 'danger')
            
        flash('Proses unggah file selesai diproses.', 'success')
        return redirect(url_for('main.kelola_panduan'))
        
    existing_templates = []
    for i in range(1, 15):
        if os.path.exists(os.path.join(panduan_dir, f'Template_Persetujuan_Tahap_{i}.docx')) or \
           os.path.exists(os.path.join(panduan_dir, f'Template_Persetujuan_Tahap_{i}.doc')):
            existing_templates.append(i)

    return render_template('main/admin_upload.html', title='Kelola Tautan & Panduan', existing_templates=existing_templates)

@bp.route('/download-template-tahap/<int:tahap>')
@login_required
def download_template_tahap(tahap):
    """
    Rute dinamis untuk mengunduh template persetujuan sesuai tahap saat ini.
    Bisa diakses oleh mahasiswa maupun dosen.
    """
    panduan_dir = os.path.join(current_app.root_path, 'static', 'panduan')
    
    # Cek format docx dulu, lalu doc
    filename_docx = f'Template_Persetujuan_Tahap_{tahap}.docx'
    filename_doc = f'Template_Persetujuan_Tahap_{tahap}.doc'
    
    if os.path.exists(os.path.join(panduan_dir, filename_docx)):
        return send_from_directory(panduan_dir, filename_docx, as_attachment=True)
    elif os.path.exists(os.path.join(panduan_dir, filename_doc)):
        return send_from_directory(panduan_dir, filename_doc, as_attachment=True)
    else:
        flash(f'Template Persetujuan untuk Tahap {tahap} belum diunggah oleh Admin.', 'warning')
        return redirect(request.referrer)
    
@bp.route('/database-viewer')
@login_required
def database_viewer():
    if current_user.role not in ['dosen', 'admin']:
        abort(403)
        
    context = {
        'users': User.query.all(),
        'lecturers': Lecturer.query.all(),
        'students': Student.query.all(),
        'supervisions': Supervision.query.all(),
        'seminars': Seminar.query.all(),
        'evaluations': Evaluation.query.all(),
        'honorariums': Honorarium.query.all(),
        'documents': Document.query.all(),
        'bimbingan_requests': Bimbingan.query.all(),
        'chat_messages': ChatMessage.query.all()
    }
    
    return render_template('main/db_viewer.html', title='Database Viewer', **context)

@bp.route('/persetujuan-pembimbing')
@login_required
def persetujuan_pembimbing():
    if current_user.role != 'dosen': abort(403)
    applications = SupervisorApplication.query.filter_by(lecturer_id=current_user.lecturer.id, status='Pending').all()
    return render_template('main/persetujuan_pembimbing.html', applications=applications)

@bp.route('/handle_pengajuan/<int:application_id>', methods=['POST'])
@login_required
def handle_pengajuan(application_id):
    if current_user.role != 'dosen': abort(403)
    
    application = SupervisorApplication.query.get_or_404(application_id)
    if application.lecturer_id != current_user.lecturer.id: abort(403)
    
    action = request.form.get('action')
    if action == 'approve':
        application.status = 'Approved'
        existing = Supervision.query.filter_by(lecturer_id=application.lecturer_id, student_id=application.student_id).first()
        if not existing:
            supervision = Supervision(
                lecturer_id=application.lecturer_id,
                student_id=application.student_id,
                role=application.role_applied_for
            )
            db.session.add(supervision)
            if application.research_title_applied and not application.student.research_title:
                 application.student.research_title = application.research_title_applied
        flash('Pengajuan disetujui.', 'success')
        
    elif action == 'reject':
        application.status = 'Rejected'
        flash('Pengajuan ditolak. Mahasiswa harus mengajukan ulang.', 'warning')
        
    db.session.commit()
    return redirect(url_for('main.persetujuan_pembimbing'))

def check_and_advance_progress(student):
    required_tasks = TAHAPAN_TUGAS.get(student.current_progress_step, [])
    if not required_tasks: return False
    
    all_tasks_approved = True
    for task in required_tasks:
        doc = Document.query.filter_by(student_id=student.id, document_type=task['type'], progress_step=student.current_progress_step).first()
        if not doc:
            all_tasks_approved = False
            break
        if doc.approvals.count() == 0:
            all_tasks_approved = False
            break
            
    if all_tasks_approved:
        student.current_progress_step += 1
        
        status_map = {
            1: "Bimbingan Tahap 1", 
            2: "Sidang Komisi Pra-Seminar Proposal", 
            3: "Sidang Seminar Proposal",
            4: "Bimbingan Tahap 2", 
            5: "Sidang Komisi Pra-Seminar Hasil", 
            6: "Sidang Seminar Hasil",
            7: "Bimbingan Tahap 3", 
            8: "Sidang Komisi Pra-Ujian Tertutup",
            9: "Sidang Tertutup", 
            10: "Perbaikan Sidang", 
            11: "Sidang Terbuka (Promosi)",
            12: "Revisi Final", 
            13: "Yudisium",
            14: "Proses Kelulusan"
        }
        
        student.status_bimbingan = status_map.get(student.current_progress_step, f"Tahap {student.current_progress_step}")
        
        if student.current_progress_step >= 14:
            student.status_bimbingan = "Lulus Program Doktoral"
            
        db.session.commit()
        return True
    return False

def advance_student_step(student):
    progress_steps = {
        1: "Bimbingan Tahap 1", 
        2: "Sidang Komisi Pra-Seminar Proposal", 
        3: "Sidang Seminar Proposal",
        4: "Bimbingan Tahap 2", 
        5: "Sidang Komisi Pra-Seminar Hasil", 
        6: "Sidang Seminar Hasil",
        7: "Bimbingan Tahap 3", 
        8: "Sidang Komisi Pra-Ujian Tertutup",
        9: "Sidang Tertutup", 
        10: "Perbaikan Sidang", 
        11: "Sidang Terbuka (Promosi)",
        12: "Revisi Final", 
        13: "Yudisium",
        14: "Proses Kelulusan"
    }
    
    seminar_creation_map = { 3: 'Sidang Komisi Pra-Seminar Proposal', 4: 'Seminar Proposal', 6: 'Sidang Komisi Pra-Seminar Hasil', 7: 'Seminar Hasil', 9: 'Sidang Komisi Pra-Ujian Tertutup', 10: 'Sidang Tertutup', 12: 'Sidang Terbuka' }

    if student.current_progress_step < 14:
        student.current_progress_step += 1
        new_step = student.current_progress_step
        student.status_bimbingan = progress_steps.get(new_step, "Selesai")

        if new_step in seminar_creation_map:
            seminar_type = seminar_creation_map[new_step]
            if not Seminar.query.filter_by(student_id=student.id, seminar_type=seminar_type).first():
                db.session.add(Seminar(student_id=student.id, seminar_type=seminar_type, date=datetime.date.today()))
        
        if new_step == 14: 
            for sup in student.supervisions:
                db.session.add(Honorarium(lecturer_id=sup.lecturer_id, student_id=student.id, activity=f"Bimbingan Kelulusan - {student.name}", amount=1500000))
        
        db.session.commit()
        return True
    return False

@bp.route('/approve_document', methods=['POST'])
@login_required
def approve_document():
    if current_user.role != 'dosen': return jsonify({'status': 'error', 'message': 'Unauthorized'}), 403
    
    data = request.get_json()
    doc_id = data.get('document_id')
    document = Document.query.get_or_404(doc_id)

    existing_approval = DocumentApproval.query.filter_by(document_id=doc_id, lecturer_id=current_user.lecturer.id).first()
    if not existing_approval:
        approval = DocumentApproval(document_id=doc_id, lecturer_id=current_user.lecturer.id, is_approved=True)
        db.session.add(approval)
        db.session.commit()
    
    advanced = check_and_advance_progress(document.student)
    
    msg = 'Dokumen berhasil disetujui.'
    if advanced: msg += ' Mahasiswa telah naik ke tahap berikutnya!'
    
    return jsonify({'status': 'success', 'message': msg, 'advanced_to_next_step': advanced})

@bp.route('/retract_document_approval', methods=['POST'])
@login_required
def retract_document_approval():
    if current_user.role != 'dosen': return jsonify({'status': 'error', 'message': 'Unauthorized'}), 403
    data = request.get_json()
    doc_id = data.get('document_id')
    
    approval = DocumentApproval.query.filter_by(document_id=doc_id, lecturer_id=current_user.lecturer.id).first()
    if approval:
        db.session.delete(approval)
        db.session.commit()
        return jsonify({'status': 'success', 'message': 'Persetujuan ditarik kembali.'})
    
    return jsonify({'status': 'error', 'message': 'Persetujuan tidak ditemukan.'}), 400

@bp.route('/upload_profile_picture', methods=['POST'])
@login_required
def upload_profile_picture():
    if 'profile_picture' not in request.files:
        return jsonify({'message': 'Tidak ada file yang diunggah'}), 400
    
    file = request.files['profile_picture']
    if file.filename == '':
        return jsonify({'message': 'Tidak ada file yang dipilih'}), 400
        
    if file and file.filename.split('.')[-1].lower() in ALLOWED_EXTENSIONS:
        user_obj = None
        prefix = ""
        if current_user.role == 'dosen':
            user_obj = current_user.lecturer
            prefix = "dosen"
        elif current_user.role == 'mahasiswa':
            user_obj = current_user.student
            prefix = "mhs"
            
        if not user_obj: return jsonify({'message': 'User profile not found'}), 404

        if user_obj.profile_picture:
            old_path = os.path.join(current_app.config['UPLOAD_FOLDER'], user_obj.profile_picture)
            if os.path.exists(old_path): os.remove(old_path)

        filename = secure_filename(f"pfp_{prefix}_{current_user.id}_{int(datetime.datetime.now().timestamp())}.{file.filename.split('.')[-1]}")
        file.save(os.path.join(current_app.config['UPLOAD_FOLDER'], filename))
        
        user_obj.profile_picture = filename
        db.session.commit()
        return jsonify({'message': 'Foto profil berhasil diperbarui'}), 200
    else:
        return jsonify({'message': 'Format file tidak diizinkan'}), 400

@bp.route('/delete_profile_picture', methods=['POST'])
@login_required
def delete_profile_picture():
    filename = None
    if current_user.role == 'mahasiswa' and current_user.student.profile_picture:
        filename = current_user.student.profile_picture
        current_user.student.profile_picture = None
    elif current_user.role == 'dosen' and current_user.lecturer.profile_picture:
        filename = current_user.lecturer.profile_picture
        current_user.lecturer.profile_picture = None

    if filename:
        filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
        if os.path.exists(filepath): os.remove(filepath)
        db.session.commit()
        return jsonify({'message': 'Foto profil berhasil dihapus'}), 200
    else:
        return jsonify({'message': 'Tidak ada foto profil untuk dihapus'}), 400
    
@bp.route('/skip_to_yudisium/<int:student_id>', methods=['POST'])
@login_required
def skip_to_yudisium(student_id):
    if current_user.role != 'dosen': abort(403)
    student = Student.query.get_or_404(student_id)
    
    student.current_progress_step = 14
    student.status_bimbingan = "Lulus Program Doktoral"
    db.session.commit()
    return jsonify({'status': 'success', 'message': 'Mahasiswa berhasil diluluskan.'})